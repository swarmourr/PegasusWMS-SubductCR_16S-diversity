"""
gaslib.py  --  shared logic for the noble-gas Pegasus jobs
==========================================================
Version-independent reproduction of the GAS_PROCESS calculated tabs. Every
wrapper in bin/ imports this. Nothing here is workflow-specific; it just
parses the workbook's own formulas, reads its constants/tables/selection, and
recomputes the calculated columns (validated to machine precision vs. the
Aug-2024 and Oct-2026 sheets).

File contracts between jobs:
  semi_raw.csv / blanks.csv / air_stds.csv / samples.csv :
        header = _row,_code,<join columns...>   (one row per run)
  blank_ref.json / airstd_ref.json :
        {"<beam col letter>": <average value>, ...}
  *_overlay.csv (blanks_tab / air_stds_tab / samples_tab) :
        header = sheet,row,col,value            (cells to overlay on assembly)
"""
import re
import csv
import json
from collections import Counter
from statistics import mean
from openpyxl import load_workbook, Workbook
from openpyxl.utils import column_index_from_string as ci, get_column_letter as gl
from openpyxl.styles import Font

ARIAL = Font(name="Arial", size=10)


# ----------------------------- small helpers ------------------------------
def num(v):
    return v if isinstance(v, (int, float)) else None

def to_int(v):
    try:
        return int(round(float(v)))
    except (TypeError, ValueError):
        return None

def code2(v):
    return str(v).split(".")[0]

def avg_rows(formula):
    """Exact rows referenced by =AVERAGE(...)  (cells + ranges, any order)."""
    body = formula.split("AVERAGE", 1)[1]
    rows = set()
    for a, b in re.findall(r"[A-Z]+(\d+):[A-Z]+(\d+)", body):
        rows.update(range(int(a), int(b) + 1))
    for r in re.findall(r"[A-Z]+(\d+)", re.sub(r"[A-Z]+\d+:[A-Z]+\d+", "", body)):
        rows.add(int(r))
    return sorted(rows)

def read_table(ws, code_col, r1, r2, val_col):
    t = {}
    for r in range(r1, r2 + 1):
        c, v = to_int(ws.cell(r, ci(code_col)).value), num(ws.cell(r, ci(val_col)).value)
        if c is not None and v is not None:
            t[c] = v
    return t

def lookup(code, table):
    """Excel LOOKUP: largest key <= code (clamps out of range)."""
    keys = sorted(table)
    chosen = keys[0]
    for k in keys:
        if k <= code:
            chosen = k
        else:
            break
    return table[chosen]

def parse_conc(formula):
    """Parse a concentration formula into its wiring (works across layouts)."""
    w = {}
    m = re.search(r"Blanks!\$?([A-Z]+)\$?(\d+)", formula)
    w["beam"], w["blank_row"] = m.group(1), int(m.group(2))
    m = re.search(r"Air STDs'?!\$?([A-Z]+)\$?(\d+)", formula)
    w["air_col"], w["air_row"] = m.group(1), int(m.group(2))
    m = re.search(r"'isotope abundances'!\$?([A-Z]+)\$?(\d+)", formula)
    if m:
        w["conc"] = ("isotope abundances", m.group(1), int(m.group(2)))
    else:
        s = re.sub(r"(Blanks|Air STDs)'?!\$?[A-Z]+\$?\d+", "", formula)
        mm = re.search(r"(?<![A-Z'])([A-Z]{1,2})\$(\d+)", s)
        w["conc"] = ("Samples", mm.group(1), int(mm.group(2)))
    w["depl"] = "0.7466" in formula
    L = re.findall(
        r"LOOKUP\(\s*(?:VALUE\()?\$?([A-Z]+)\d+\)?,"
        r"\$?([A-Z]+)\$?(\d+):\$?[A-Z]+\$?(\d+),\$?([A-Z]+)\$?(\d+):\$?[A-Z]+\$?(\d+)\)", formula)
    w["split"] = (L[0][0], L[0][1], int(L[0][2]), int(L[0][3]), L[0][4], int(L[0][5]), int(L[0][6]))
    w["line"] = (L[1][0], L[1][1], int(L[1][2]), int(L[1][3]), L[1][4], int(L[1][5]), int(L[1][6]))
    w["man"] = re.search(r"/\$?([A-Z]+)\d+\s*$", formula.strip()).group(1)
    return w


# ============================ workbook context ============================
class Context:
    """Loads the workbook and exposes the shared data every job needs."""

    def __init__(self, wb_path):
        self.wf = load_workbook(wb_path, data_only=False)   # formulas
        self.wv = load_workbook(wb_path, data_only=True)     # cached values
        self.SR = self.wv["Semi Raw"]
        self._blank_avg = None
        self._air_avg = None
        self._build()

    def locate(self, sheet):
        ws = self.wf[sheet]
        hdr = next((r for r in range(1, 14)
                    if any(str(ws.cell(r, c).value) == "Date" for c in range(1, 6))), 1)
        return hdr, hdr + 1

    def _build(self):
        SRf = self.wf["Semi Raw"]
        hdr, drow = self.locate("Semi Raw")
        self.join = {}
        for c in range(1, SRf.max_column + 1):
            f = SRf.cell(drow, c).value
            if isinstance(f, str) and f.startswith("="):
                m = re.match(r"='?([A-Za-z0-9 ]+?)'?!\$?([A-Z]+)\d+$", f)
                if m:
                    self.join[gl(c)] = (m.group(1), m.group(2))
        self.join_cols = list(self.join)
        self.beam_cols = [c for c, (sh, _) in self.join.items()
                          if "Paste" in sh and sh != "SampleList Paste"]
        self.rows = []
        for r in range(drow, self.SR.max_row + 1):
            if self.SR.cell(r, ci("L")).value in (None, ""):
                continue
            row = {"_row": r, "_code": code2(self.SR.cell(r, ci("L")).value)}
            for col in self.join:
                row[col] = self.SR.cell(r, ci(col)).value
            self.rows.append(row)
        self.blanks = [r for r in self.rows if r["_code"] == "2"]
        self.airs = [r for r in self.rows if r["_code"] == "1"]
        self.samps = [r for r in self.rows if r["_code"] == "3"]
        self.SIG = [c for c in ["O", "P", "Q", "U", "AD", "AL"] if c in self.beam_cols] \
            or self.beam_cols[:6]
        self.bidx = {self._sig(r): r for r in self.blanks}
        self.aidx = {self._sig(r): r for r in self.airs}

    def _sig(self, row):
        return tuple(round(num(row[c]), 4) if num(row[c]) is not None else None for c in self.SIG)

    def _sig_ws(self, ws, r):
        return tuple(round(num(ws.cell(r, ci(c)).value), 4)
                     if num(ws.cell(r, ci(c)).value) is not None else None for c in self.SIG)

    def sel_rows(self, sheet, avg_formula):
        idx = self.bidx if sheet == "Blanks" else self.aidx
        ws = self.wv[sheet]
        return [idx[self._sig_ws(ws, r)] for r in avg_rows(avg_formula) if self._sig_ws(ws, r) in idx]

    def _avg_row(self, sheet, label="avg"):
        ws = self.wv[sheet]
        return next((r for r in range(1, ws.max_row + 1)
                     if str(ws.cell(r, ci("A")).value).strip().lower() == label), None)

    def blank_averages(self):
        if self._blank_avg is None:
            r = self._avg_row("Blanks")
            sel = self.sel_rows("Blanks", self.wf["Blanks"].cell(r, ci(self.beam_cols[0])).value) if r else []
            self.blank_avg_row = r
            self._blank_avg = {c: mean(v for v in (num(x[c]) for x in sel) if v is not None)
                               for c in self.beam_cols if any(num(x[c]) is not None for x in sel)}
        return self._blank_avg

    def air_averages(self):
        if self._air_avg is None:
            r = self._avg_row("Air STDs")
            sel = self.sel_rows("Air STDs", self.wf["Air STDs"].cell(r, ci(self.beam_cols[0])).value) if r else []
            self.air_avg_row = r
            self._air_avg = {c: mean(v for v in (num(x[c]) for x in sel) if v is not None)
                             for c in self.beam_cols if any(num(x[c]) is not None for x in sel)}
        return self._air_avg

    def copy_all(self, out):
        for name in self.wv.sheetnames:
            src = self.wv[name]
            ws = out.create_sheet(name)
            for row in src.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        ws.cell(cell.row, cell.column, cell.value).font = ARIAL


# ========================= per-tab overlay builders =======================
def blanks_overlay(ctx, blank_avg):
    r = ctx._avg_row("Blanks")
    return [("Blanks", r, c, v) for c, v in blank_avg.items()] if r else []


def air_stds_overlay(ctx, blank_avg):
    src, srcv = ctx.wf["Air STDs"], ctx.wv["Air STDs"]
    hdr, drow = ctx.locate("Air STDs")
    bc, ratio = {}, {}
    for c in range(1, src.max_column + 1):
        f = str(src.cell(drow, c).value).replace(" ", "")
        m = re.fullmatch(r"=\$?([A-Z]+)\d+-Blanks!\$?([A-Z]+)\$?\d+", f)
        if m:
            bc[gl(c)] = m.group(1)
            continue
        m = re.fullmatch(r"=\(?([A-Z]+)\d+/([A-Z]+)\d+\)?", f)
        if m:
            ratio[gl(c)] = (m.group(1), m.group(2))
    overlay = []
    for i, row in enumerate(ctx.airs):
        rr = drow + i
        cell = {}
        for col, raw in bc.items():
            rv = num(row.get(raw))
            v = (rv - blank_avg[raw]) if (rv is not None and raw in blank_avg) else rv
            cell[col] = v
            if v is not None:
                overlay.append(("Air STDs", rr, col, v))
        for col, (a, b) in ratio.items():
            va, vb = cell.get(a), cell.get(b)
            v = va / vb if (va is not None and vb) else None
            cell[col] = v
            if v is not None:
                overlay.append(("Air STDs", rr, col, v))
    return overlay


def samples_overlay(ctx, blank_avg, air_avg):
    src, srcv = ctx.wf["Samples"], ctx.wv["Samples"]
    hdr, drow = ctx.locate("Samples")
    conc_cols, ratio_cols = [], {}
    for c in range(1, src.max_column + 1):
        h = str(src.cell(hdr, c).value)
        f0 = str(src.cell(drow, c).value)
        if "cm3/cm3STP" in h and "Blanks!" in f0:
            conc_cols.append(gl(c))
        elif "/" in h and "cm3" not in h:
            m = re.fullmatch(r"=\(?([A-Z]+)\d+/([A-Z]+)\d+\)?(?:\*([0-9.]+))?", f0.replace(" ", ""))
            if m:
                ratio_cols[gl(c)] = (m.group(1), m.group(2), float(m.group(3)) if m.group(3) else 1.0)
    wiring = {}
    for c in conc_cols:
        cand = Counter()
        for rr in range(drow, min(drow + 40, drow + len(ctx.samps))):
            f = src.cell(rr, ci(c)).value
            if isinstance(f, str) and "Blanks!" in f:
                try:
                    cand[repr(parse_conc(f))] += 1
                except Exception:
                    pass
        if cand:
            wiring[c] = eval(cand.most_common(1)[0][0])

    overlay = []
    results = {drow + i: {} for i in range(len(ctx.samps))}
    for c, w in wiring.items():
        beam = w["beam"]
        b, a = blank_avg.get(beam), air_avg.get(beam)
        cs, cc, cr = w["conc"]
        conc_val = num(ctx.wv[cs].cell(cr, ci(cc)).value)
        if b is None or a is None or conc_val is None or a == b:
            continue
        sp = read_table(srcv, w["split"][1], w["split"][2], w["split"][3], w["split"][4])
        ln = read_table(srcv, w["line"][1], w["line"][2], w["line"][3], w["line"][4])
        for i, row in enumerate(ctx.samps):
            rr = drow + i
            bs = num(row.get(beam))
            man = num(srcv.cell(rr, ci(w["man"])).value)
            csp, cln = to_int(row.get(w["split"][0])), to_int(row.get(w["line"][0]))
            if None in (bs, man, csp, cln) or man == 0:
                continue
            depl = (1 - 0.7466 / 1990) ** num(row.get("E")) if w["depl"] else 1.0
            v = (bs - b) / (a - b) * conc_val * depl * (lookup(csp, sp) / lookup(cln, ln)) / man
            results[rr][c] = v
            overlay.append(("Samples", rr, c, v))
    for rc, (a, b, k) in ratio_cols.items():
        if a in conc_cols and b in conc_cols:
            for rr, row in results.items():
                if a in row and b in row and row[b]:
                    overlay.append(("Samples", rr, rc, row[a] / row[b] * k))
    return overlay


# ============================== file I/O ==================================
def write_rows_csv(path, rows, join_cols):
    cols = ["_row", "_code"] + join_cols
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(cols)
        for r in rows:
            w.writerow([r.get(c) for c in cols])

def write_json(path, obj):
    with open(path, "w") as fh:
        json.dump(obj, fh)

def read_json(path):
    with open(path) as fh:
        return json.load(fh)

def write_overlay_csv(path, overlay):
    with open(path, "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["sheet", "row", "col", "value"])
        for sheet, row, col, value in overlay:
            w.writerow([sheet, row, col, value])

def read_overlay_csv(path):
    out = []
    with open(path) as fh:
        r = csv.DictReader(fh)
        for d in r:
            out.append((d["sheet"], int(d["row"]), d["col"], float(d["value"])))
    return out
