"""
noble_gas_reducer.py  --  OOP noble-gas tab reproducer
=======================================================
Copies the whole workbook through (nothing is ever dropped) and overlays
Python-reproduced values on each calculated tab. One class per table:

    SemiRawTab   validates the paste -> Semi Raw join (keeps manual overrides)
    BlanksTab    reproduces the 'avg' row
    AirStdsTab   reproduces blank-corrected columns + ratio columns + AVG row
    SamplesTab   reproduces concentration columns + simple ratio columns

Everything (columns, constants, lookup tables, blank/air-std selection,
depletion on/off) is auto-detected from the file, so the same code works
across sheet revisions (Aug-2024, Oct-2026, ...).

    python src/noble_gas_reducer.py
    python src/noble_gas_reducer.py  <input_workbook.xlsx>  [output.xlsx]
    python src/noble_gas_reducer.py  <input_workbook.xlsx>  <prefix>  --separate
        (--separate writes one .xlsx PER TAB, named <prefix>__<TabName>.xlsx)
"""
import sys
import os
import re
from collections import Counter
from statistics import mean
from openpyxl import load_workbook, Workbook
from openpyxl.utils import column_index_from_string as ci, get_column_letter as gl
from openpyxl.styles import Font

ARIAL = Font(name="Arial", size=10)


# ----------------------------- small helpers ------------------------------
def num(v):
    return v if isinstance(v, (int, float)) else None

def to_int(v):
    try:
        return int(round(float(v)))
    except (TypeError, ValueError):
        return None

def code2(v):
    return str(v).split(".")[0]

def avg_rows(formula):
    """Exact rows referenced by =AVERAGE(...)  (cells + ranges, any order)."""
    body = formula.split("AVERAGE", 1)[1]
    rows = set()
    for a, b in re.findall(r"[A-Z]+(\d+):[A-Z]+(\d+)", body):
        rows.update(range(int(a), int(b) + 1))
    for r in re.findall(r"[A-Z]+(\d+)", re.sub(r"[A-Z]+\d+:[A-Z]+\d+", "", body)):
        rows.add(int(r))
    return sorted(rows)

def read_table(ws, code_col, r1, r2, val_col):
    t = {}
    for r in range(r1, r2 + 1):
        c, v = to_int(ws.cell(r, ci(code_col)).value), num(ws.cell(r, ci(val_col)).value)
        if c is not None and v is not None:
            t[c] = v
    return t

def lookup(code, table):
    """Excel LOOKUP: largest key <= code (clamps out of range)."""
    keys = sorted(table)
    chosen = keys[0]
    for k in keys:
        if k <= code:
            chosen = k
        else:
            break
    return table[chosen]

def parse_conc(formula):
    """Parse a concentration formula into its wiring (works across layouts)."""
    w = {}
    m = re.search(r"Blanks!\$?([A-Z]+)\$?(\d+)", formula)
    w["beam"], w["blank_row"] = m.group(1), int(m.group(2))
    m = re.search(r"Air STDs'?!\$?([A-Z]+)\$?(\d+)", formula)
    w["air_col"], w["air_row"] = m.group(1), int(m.group(2))
    m = re.search(r"'isotope abundances'!\$?([A-Z]+)\$?(\d+)", formula)
    if m:
        w["conc"] = ("isotope abundances", m.group(1), int(m.group(2)))
    else:
        s = re.sub(r"(Blanks|Air STDs)'?!\$?[A-Z]+\$?\d+", "", formula)
        mm = re.search(r"(?<![A-Z'])([A-Z]{1,2})\$(\d+)", s)
        w["conc"] = ("Samples", mm.group(1), int(mm.group(2)))
    w["depl"] = "0.7466" in formula
    L = re.findall(
        r"LOOKUP\(\s*(?:VALUE\()?\$?([A-Z]+)\d+\)?,"
        r"\$?([A-Z]+)\$?(\d+):\$?[A-Z]+\$?(\d+),\$?([A-Z]+)\$?(\d+):\$?[A-Z]+\$?(\d+)\)", formula)
    w["split"] = (L[0][0], L[0][1], int(L[0][2]), int(L[0][3]), L[0][4], int(L[0][5]), int(L[0][6]))
    w["line"] = (L[1][0], L[1][1], int(L[1][2]), int(L[1][3]), L[1][4], int(L[1][5]), int(L[1][6]))
    w["man"] = re.search(r"/\$?([A-Z]+)\d+\s*$", formula.strip()).group(1)
    return w


# ============================= workbook context ============================
class GasWorkbook:
    """Loads the workbook and holds data shared by all tab reproducers."""

    def __init__(self, path):
        self.wf = load_workbook(path, data_only=False)   # formulas
        self.wv = load_workbook(path, data_only=True)    # cached values
        self.SR = self.wv["Semi Raw"]
        self._blank_avg = None
        self._air_avg = None
        self._build_context()

    # -- locate header / first data row of a sheet --
    def locate(self, sheet):
        ws = self.wf[sheet]
        hdr = next((r for r in range(1, 14)
                    if any(str(ws.cell(r, c).value) == "Date" for c in range(1, 6))), 1)
        return hdr, hdr + 1

    # -- Semi Raw join map + rows (read Semi Raw values, incl. manual overrides) --
    def _build_context(self):
        SRf = self.wf["Semi Raw"]
        hdr, drow = self.locate("Semi Raw")
        self.join = {}
        for c in range(1, SRf.max_column + 1):
            f = SRf.cell(drow, c).value
            if isinstance(f, str) and f.startswith("="):
                m = re.match(r"='?([A-Za-z0-9 ]+?)'?!\$?([A-Z]+)\d+$", f)
                if m:
                    self.join[gl(c)] = (m.group(1), m.group(2))
        self.beam_cols = [c for c, (sh, _) in self.join.items()
                          if "Paste" in sh and sh != "SampleList Paste"]
        self.rows = []
        for r in range(drow, self.SR.max_row + 1):
            if self.SR.cell(r, ci("L")).value in (None, ""):
                continue
            row = {"_row": r}
            for col in self.join:
                row[col] = self.SR.cell(r, ci(col)).value
            self.rows.append(row)
        self.blanks = [r for r in self.rows if code2(r.get("L")) == "2"]
        self.airs = [r for r in self.rows if code2(r.get("L")) == "1"]
        self.samps = [r for r in self.rows if code2(r.get("L")) == "3"]
        self.SIG = [c for c in ["O", "P", "Q", "U", "AD", "AL"] if c in self.beam_cols] \
            or self.beam_cols[:6]
        self.bidx = {self._sig(r): r for r in self.blanks}
        self.aidx = {self._sig(r): r for r in self.airs}

    def _sig(self, row):
        return tuple(round(num(row[c]), 4) if num(row[c]) is not None else None for c in self.SIG)

    def _sig_ws(self, ws, r):
        return tuple(round(num(ws.cell(r, ci(c)).value), 4)
                     if num(ws.cell(r, ci(c)).value) is not None else None for c in self.SIG)

    def sel_rows(self, sheet, avg_formula):
        idx = self.bidx if sheet == "Blanks" else self.aidx
        ws = self.wv[sheet]
        return [idx[self._sig_ws(ws, r)] for r in avg_rows(avg_formula) if self._sig_ws(ws, r) in idx]

    def _avg_row(self, sheet, label):
        ws = self.wv[sheet]
        return next((r for r in range(1, ws.max_row + 1)
                     if str(ws.cell(r, ci("A")).value).strip().lower() == label), None)

    # -- shared references (computed once) --
    def blank_averages(self):
        if self._blank_avg is None:
            r = self._avg_row("Blanks", "avg")
            sel = self.sel_rows("Blanks", self.wf["Blanks"].cell(r, ci(self.beam_cols[0])).value) if r else []
            self._blank_avg_row = r
            self._blank_avg = {c: mean(v for v in (num(x[c]) for x in sel) if v is not None)
                               for c in self.beam_cols if any(num(x[c]) is not None for x in sel)}
        return self._blank_avg

    def air_averages(self):
        if self._air_avg is None:
            r = self._avg_row("Air STDs", "avg")
            sel = self.sel_rows("Air STDs", self.wf["Air STDs"].cell(r, ci(self.beam_cols[0])).value) if r else []
            self._air_avg_row = r
            self._air_avg = {c: mean(v for v in (num(x[c]) for x in sel) if v is not None)
                             for c in self.beam_cols if any(num(x[c]) is not None for x in sel)}
        return self._air_avg

    def copy_all(self, out):
        """Copy every tab (values) so nothing is ever dropped."""
        for name in self.wv.sheetnames:
            src = self.wv[name]
            ws = out.create_sheet(name)
            for row in src.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        ws.cell(cell.row, cell.column, cell.value).font = ARIAL

    def reproduce(self, out):
        self.copy_all(out)
        self.tabs = [SemiRawTab(self), BlanksTab(self), AirStdsTab(self), SamplesTab(self)]
        for tab in self.tabs:
            tab.reproduce(out)
        return self.tabs


# ============================= tab reproducers =============================
class Tab:
    name = ""

    def __init__(self, wb):
        self.wb = wb
        self.report = []          # (col, n_cells, worst_rel_err)
        self.note = ""

    def reproduce(self, out):
        raise NotImplementedError

    def log(self, col, n, worst):
        self.report.append((col, n, worst))

    def summary(self):
        cols = len(self.report)
        cells = sum(n for _, n, _ in self.report)
        worst = max((w for _, _, w in self.report if w < 1), default=0.0)
        anom = [(c, w) for c, _, w in self.report if w >= 1e-6]
        return cols, cells, worst, anom


class SemiRawTab(Tab):
    """Semi Raw is the paste join; validate it and keep any manual overrides."""
    name = "Semi Raw"

    def reproduce(self, out):
        wb = self.wb
        srcs = {s: wb.wv[s] for s, _ in wb.join.values()}
        overrides = 0
        for col in wb.beam_cols:
            sh, sc = wb.join[col]
            worst, n = 0.0, 0
            for row in wb.rows:
                paste = num(srcs[sh].cell(row["_row"], ci(sc)).value)
                semi = num(row[col])
                if paste is None or semi is None:
                    continue
                e = abs(paste - semi) / max(abs(semi), 1e-30)
                if e > 1e-9:
                    overrides += 1
                else:
                    worst = max(worst, e)
                n += 1
            self.log(col, n, worst)
        self.note = f"{overrides} manual override cell(s) kept" if overrides else ""


class BlanksTab(Tab):
    """Reproduce the 'avg' row from the selected blank runs."""
    name = "Blanks"

    def reproduce(self, out):
        wb = self.wb
        ws, srcv = out["Blanks"], wb.wv["Blanks"]
        bavg = wb.blank_averages()
        r = wb._blank_avg_row
        if not r:
            return
        for c, v in bavg.items():
            ws.cell(r, ci(c), v).font = ARIAL
            exp = num(srcv.cell(r, ci(c)).value)
            if exp is not None:
                self.log(c, 1, abs(v - exp) / max(abs(exp), 1e-30))


class AirStdsTab(Tab):
    """Reproduce blank-corrected columns, ratio columns, and the AVG row."""
    name = "Air STDs"

    def reproduce(self, out):
        wb = self.wb
        ws, src, srcv = out["Air STDs"], wb.wf["Air STDs"], wb.wv["Air STDs"]
        hdr, drow = wb.locate("Air STDs")
        bavg = wb.blank_averages()

        bc, ratio = {}, {}
        for c in range(1, src.max_column + 1):
            f = str(src.cell(drow, c).value).replace(" ", "")
            m = re.fullmatch(r"=\$?([A-Z]+)\d+-Blanks!\$?([A-Z]+)\$?\d+", f)
            if m:
                bc[gl(c)] = m.group(1)
                continue
            m = re.fullmatch(r"=\(?([A-Z]+)\d+/([A-Z]+)\d+\)?", f)
            if m:
                ratio[gl(c)] = (m.group(1), m.group(2))

        percol = {c: {} for c in list(bc) + list(ratio)}
        for i, row in enumerate(wb.airs):
            rr = drow + i
            cell = {}
            for col, raw in bc.items():
                rv = num(row.get(raw))
                v = (rv - bavg[raw]) if (rv is not None and raw in bavg) else rv
                cell[col] = v
                if v is not None:
                    ws.cell(rr, ci(col), v).font = ARIAL
                    percol[col][rr] = v
            for col, (a, b) in ratio.items():
                va, vb = cell.get(a), cell.get(b)
                v = va / vb if (va is not None and vb) else None
                cell[col] = v
                if v is not None:
                    ws.cell(rr, ci(col), v).font = ARIAL
                    percol[col][rr] = v

        for col in list(bc) + list(ratio):
            worst, n = 0.0, 0
            for i in range(len(wb.airs)):
                rr = drow + i
                got, exp = percol[col].get(rr), num(srcv.cell(rr, ci(col)).value)
                if got is None or exp is None:
                    continue
                worst = max(worst, abs(got - exp) / max(abs(exp), 1e-30))
                n += 1
            if n:
                self.log(col, n, worst)

        # AVG row: raw beam averages
        r = next((rr for rr in range(1, srcv.max_row + 1)
                  if str(srcv.cell(rr, ci("A")).value).strip().upper() == "AVG"), None)
        if r:
            for c, v in wb.air_averages().items():
                ws.cell(r, ci(c), v).font = ARIAL


class SamplesTab(Tab):
    """Reproduce concentration columns and simple ratio columns."""
    name = "Samples"

    def reproduce(self, out):
        wb = self.wb
        ws, src, srcv = out["Samples"], wb.wf["Samples"], wb.wv["Samples"]
        hdr, drow = wb.locate("Samples")
        bavg, aavg = wb.blank_averages(), wb.air_averages()

        conc_cols, ratio_cols = [], {}
        for c in range(1, src.max_column + 1):
            h = str(src.cell(hdr, c).value)
            f0 = str(src.cell(drow, c).value)
            if "cm3/cm3STP" in h and "Blanks!" in f0:
                conc_cols.append(gl(c))
            elif "/" in h and "cm3" not in h:
                m = re.fullmatch(r"=\(?([A-Z]+)\d+/([A-Z]+)\d+\)?(?:\*([0-9.]+))?", f0.replace(" ", ""))
                if m:
                    ratio_cols[gl(c)] = (m.group(1), m.group(2), float(m.group(3)) if m.group(3) else 1.0)

        wiring = {}
        for c in conc_cols:
            cand = Counter()
            for rr in range(drow, min(drow + 40, drow + len(wb.samps))):
                f = src.cell(rr, ci(c)).value
                if isinstance(f, str) and "Blanks!" in f:
                    try:
                        cand[repr(parse_conc(f))] += 1
                    except Exception:
                        pass
            if cand:
                wiring[c] = eval(cand.most_common(1)[0][0])

        results = {drow + i: {} for i in range(len(wb.samps))}
        for c, w in wiring.items():
            beam = w["beam"]
            b, a = bavg.get(beam), aavg.get(beam)
            cs, cc, cr = w["conc"]
            conc_val = num(wb.wv[cs].cell(cr, ci(cc)).value)
            if b is None or a is None or conc_val is None or a == b:
                continue
            sp = read_table(srcv, w["split"][1], w["split"][2], w["split"][3], w["split"][4])
            ln = read_table(srcv, w["line"][1], w["line"][2], w["line"][3], w["line"][4])
            worst, n = 0.0, 0
            for i, row in enumerate(wb.samps):
                rr = drow + i
                bs = num(row.get(beam))
                man = num(srcv.cell(rr, ci(w["man"])).value)
                csp, cln = to_int(row.get(w["split"][0])), to_int(row.get(w["line"][0]))
                if None in (bs, man, csp, cln) or man == 0:
                    continue
                depl = (1 - 0.7466 / 1990) ** num(row.get("E")) if w["depl"] else 1.0
                v = (bs - b) / (a - b) * conc_val * depl * (lookup(csp, sp) / lookup(cln, ln)) / man
                results[rr][c] = v
                ws.cell(rr, ci(c), v).font = ARIAL
                exp = num(srcv.cell(rr, ci(c)).value)
                if exp is not None:
                    worst = max(worst, abs(v - exp) / max(abs(exp), 1e-30))
                    n += 1
            if n:
                self.log(c, n, worst)

        for rc, (a, b, k) in ratio_cols.items():
            if a in conc_cols and b in conc_cols:
                worst, n = 0.0, 0
                for rr, row in results.items():
                    if a in row and b in row and row[b]:
                        v = row[a] / row[b] * k
                        ws.cell(rr, ci(rc), v).font = ARIAL
                        exp = num(srcv.cell(rr, ci(rc)).value)
                        if exp is not None:
                            worst = max(worst, abs(v - exp) / max(abs(exp), 1e-30))
                            n += 1
                if n:
                    self.log(rc, n, worst)


def export_separate(out, prefix):
    """Write each tab of `out` to its own .xlsx file; return the paths."""
    base = prefix[:-5] if prefix.lower().endswith(".xlsx") else prefix
    if base.endswith(os.sep) or os.path.isdir(base):
        os.makedirs(base, exist_ok=True)
        joiner = os.path.join(base, "{}.xlsx")
    else:
        joiner = base + "__{}.xlsx"
    paths = []
    for name in out.sheetnames:
        wb1 = Workbook()
        ws1 = wb1.active
        ws1.title = name[:31]
        for row in out[name].iter_rows():
            for cell in row:
                if cell.value is not None:
                    ws1.cell(cell.row, cell.column, cell.value).font = ARIAL
        path = joiner.format(name.replace(" ", "_").replace("/", "-"))
        wb1.save(path)
        paths.append(path)
    return paths


# ================================== main ==================================
def main():
    argv = sys.argv[1:]
    separate = any(a in ("-s", "--separate", "--split") for a in argv)
    pos = [a for a in argv if not a.startswith("-")]
    in_path = pos[0] if pos else os.path.join("data", "Oct2026_GAS_PROCESS_source.xlsx")
    out_path = pos[1] if len(pos) > 1 else (
        os.path.join("outputs", "tabs") if separate
        else os.path.join("outputs", "oct_tabs_generated.xlsx")
    )

    wb = GasWorkbook(in_path)
    out = Workbook()
    out.remove(out.active)
    tabs = wb.reproduce(out)

    if separate:
        paths = export_separate(out, out_path)
        print(f"input : {in_path}")
        print(f"output: {len(paths)} separate files:")
        for p in paths:
            print(f"        {p}")
        print()
    else:
        out_dir = os.path.dirname(out_path)
        if out_dir:
            os.makedirs(out_dir, exist_ok=True)
        out.save(out_path)
        print(f"input : {in_path}")
        print(f"output: {out_path}  (all {len(out.sheetnames)} tabs kept)\n")
    print(f"  {'tab':10s} {'cols':>5s} {'cells':>8s} {'worst rel.err':>14s}   note")
    anomalies = []
    for tab in tabs:
        cols, cells, worst, anom = tab.summary()
        print(f"  {tab.name:10s} {cols:5d} {cells:8d} {worst:14.1e}   {tab.note}")
        anomalies += [(tab.name, c, w) for c, w in anom]
    if anomalies:
        print("\n  per-row sheet anomalies (nonstandard column in the sheet itself):")
        for t, c, w in anomalies:
            print(f"    {t} {c}: worst {w:.1e}")
    print("\nReproduced columns are recomputed in Python and overlaid; every other\n"
          "cell and tab (isotope abundances, lookup tables, ...) is copied through.")


if __name__ == "__main__":
    main()
