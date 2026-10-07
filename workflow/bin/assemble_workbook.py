#!/usr/bin/env python3
"""assemble_workbook: copy the workbook through + overlay every reproduced tab
   -> GAS_PROCESS_reproduced.xlsx  (one file, all tabs)"""
import sys
from openpyxl import Workbook
from openpyxl.utils import column_index_from_string as ci
from gaslib import Context, read_overlay_csv, ARIAL

def main(workbook, out_path="GAS_PROCESS_reproduced.xlsx"):
    ctx = Context(workbook)
    out = Workbook()
    out.remove(out.active)
    ctx.copy_all(out)                       # every tab + reference data, as values
    for ov_file in ["blanks_tab.csv", "air_stds_tab.csv", "samples_tab.csv"]:
        for sheet, row, col, value in read_overlay_csv(ov_file):
            out[sheet].cell(row, ci(col), value).font = ARIAL
    out.save(out_path)
    print(f"wrote {out_path}  (tabs: {', '.join(out.sheetnames)})")

if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "GAS_PROCESS.xlsx")
